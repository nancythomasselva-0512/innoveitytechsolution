import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_document():
    doc = Document()

    # Page Margins (1 inch everywhere)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("INNOVEITY TECH SOLUTION  |  EXECUTIVE OVERVIEW")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184) # Light Slate
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential  •  © 2026 Innoveity Tech Solution Private Limited  •  Standard 12 pt Executive Documentation")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(12)
    font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Color Palette definitions
    NAVY = RGBColor(15, 23, 42)      # #0F172A
    BLUE_ACCENT = RGBColor(2, 132, 199) # #0284C7
    BODY_COLOR = RGBColor(51, 65, 85) # #334155
    MUTED_COLOR = RGBColor(100, 116, 139) # #64748B

    # Helper functions for adding content
    def add_super_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text.upper())
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = BLUE_ACCENT
        return p

    def add_main_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_sub_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(18)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.italic = True
        run.font.color.rgb = MUTED_COLOR
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = BLUE_ACCENT
        return p

    def add_body_p(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = 'Calibri'
            run_b.font.size = Pt(12)
            run_b.font.bold = True
            run_b.font.color.rgb = NAVY
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.color.rgb = BODY_COLOR
        return p

    def add_bullet(bold_title, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run_b = p.add_run(bold_title + " ")
        run_b.font.name = 'Calibri'
        run_b.font.size = Pt(12)
        run_b.font.bold = True
        run_b.font.color.rgb = NAVY
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.color.rgb = BODY_COLOR
        return p

    def add_num_bullet(num_str, bold_title, text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        run_n = p.add_run(num_str + " ")
        run_n.font.name = 'Segoe UI Semibold'
        run_n.font.size = Pt(12)
        run_n.font.bold = True
        run_n.font.color.rgb = BLUE_ACCENT
        run_b = p.add_run(bold_title + " ")
        run_b.font.name = 'Calibri'
        run_b.font.size = Pt(12)
        run_b.font.bold = True
        run_b.font.color.rgb = NAVY
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.color.rgb = BODY_COLOR
        return p

    # ========================== DOCUMENT HEADER ==========================
    add_super_title("Enterprise Product Architecture & Systems Overview")
    add_main_title("INNOVEITY TECH SOLUTION")
    add_sub_title("Executive Product Overview, Technical Architecture & Institutional Operations Specification")

    # ========================== METADATA TABLE ==========================
    meta_table = doc.add_table(rows=7, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="CBD5E1", sz="6", val="single")

    metadata_rows = [
        ("Official URL Link", "https://innoveitytech.com"),
        ("Project Name", "Innoveity Tech Solution (Corporate Web & Headless CMS Platform)"),
        ("Lead Designer & Architect", "Nancy Narmadha T"),
        ("Organization & Location", "Innoveity Tech Solution Private Limited, MCC MRF Innovation Park, East Tambaram, Chennai - 600059, Tamil Nadu, India"),
        ("Technology Stack", "React 19, Vite 8, React Router v7, Framer Motion, Vanilla CSS (Design Tokens), Supabase Cloud PostgreSQL, MySQL 8 Server, Express.js REST API"),
        ("Document Scope", "Official Academic, Institutional, Enterprise & Commercial Product Overview"),
        ("Document Standard", "Formatted for Standard 12 pt Executive Documentation & Technical Specifications")
    ]

    col_widths = [Inches(2.2), Inches(4.3)]

    for idx, (label, val) in enumerate(metadata_rows):
        row = meta_table.rows[idx]
        
        # Col 0: Label
        c0 = row.cells[0]
        c0.width = col_widths[0]
        set_cell_shading(c0, "F8FAFC")
        set_cell_margins(c0, top=90, bottom=90, left=140, right=140)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = 'Segoe UI Semibold'
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = NAVY
        
        # Col 1: Value
        c1 = row.cells[1]
        c1.width = col_widths[1]
        set_cell_shading(c1, "FFFFFF")
        set_cell_margins(c1, top=90, bottom=90, left=140, right=140)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(11)
        r1.font.color.rgb = BODY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ========================== 1. PRODUCT DESCRIPTION ==========================
    add_h1("1. PRODUCT DESCRIPTION")
    add_h2("What is the Innoveity Tech Solution product, and what is its purpose?")
    add_body_p(
        "Innoveity Tech Solution is an advanced, high-performance enterprise web application, corporate digital showcase, "
        "and real-time Headless Content Management System (CMS). The product is specifically engineered to resolve "
        "the operational friction, high latency, and security vulnerabilities associated with legacy content management systems, "
        "while simultaneously operating as an integrated digital bridge between enterprise software engineering and higher education workforce development."
    )
    add_body_p(
        "Unlike static brochure websites or rigid template-based builders, Innoveity Tech Solution operates as an active, "
        "multi-tenant digital ecosystem. It delivers a public-facing corporate presence with cinematic user experiences alongside a zero-code "
        "administrative portal capable of pushing sub-second live database mutations across active client sessions worldwide."
    )

    add_body_p("Core Purpose & Strategic Objectives:", bold_prefix="Strategic Pillars: ")
    add_num_bullet("1.", "Dynamic Enterprise Digital Presence & Service Hub:",
                   "Presents an authoritative, interactive gateway articulating Innoveity's cross-disciplinary engineering capabilities spanning Artificial Intelligence, Cloud Infrastructure, Custom SaaS Platforms, Mobile App Engineering, and Enterprise Systems.")
    add_num_bullet("2.", "Real-Time Zero-Rebuild Content Management (CMS):",
                   "Empowers non-technical marketing, human resources, and business operations teams to govern live website copy, reorder team hierarchies, curate showcase case studies, upload press media, and configure route-level SEO metadata in real time without developer intervention or continuous deployment (CI/CD) rebuild cycles.")
    add_num_bullet("3.", "Resilient Dual-Sync High-Availability Database Engine:",
                   "Implements an active-passive failover data architecture integrating Supabase Cloud PostgreSQL (driven by low-latency WebSocket subscription listeners) with a self-hosted MySQL 8 relational database and automated JSON seed hydration, ensuring 99.99% data persistence and immediate multi-client synchronicity.")
    add_num_bullet("4.", "End-to-End Industry-Academia Workforce Acceleration:",
                   "Directly anchors and scales Innoveity’s proven institutional development initiatives across Tamil Nadu and South India. The platform supports an established academic track record featuring an 85% average placement improvement across 50+ partner engineering colleges and over 2,000+ students trained annually in Chennai, Coimbatore, Madurai, and surrounding regions.")
    add_num_bullet("5.", "Direct High-Conversion Client Communication Pipeline:",
                   "Provides a multi-channel inquiry dispatch system incorporating client-side EmailJS browser-to-inbox routing backed by dedicated SMTP server relay fallbacks, guaranteeing zero dropped corporate leads or institutional consultation bookings.")

    # ========================== 2. DESIGNED & DEVELOPED BY ==========================
    add_h1("2. DESIGNED & DEVELOPED BY")
    add_h2("Who designed and developed the platform?")
    add_body_p(
        "The Innoveity Tech Solution web application and CMS platform was conceptualized, designed, and engineered end-to-end "
        "by Nancy Narmadha T as an enterprise-grade digital product and scalable administrative governance platform.",
        bold_prefix="Lead Architect, Designer & Full-Stack Developer: NANCY NARMADHA T\n"
    )

    add_body_p("Key Engineering, Architectural & Design Contributions by Nancy Narmadha T:", bold_prefix="Technical Contributions: ")
    add_bullet("• Full-Stack SPA Architecture & Modern Build Pipeline:",
               "Architected the platform utilizing React 19 and Vite 8, leveraging modern ECMAScript standards, code-splitting, tree-shaking, and optimized asset pipelines that achieve near-instant initial paint times and sub-second page transitions across desktop and mobile form factors.")
    add_bullet("• Bespoke Visual Identity & Mathematical Kinetic Physics:",
               "Designed and implemented an ultra-premium visual design system utilizing custom HSL color tokens, dark-mode glassmorphic layers (#121212 / #0a0a0c), subtle backdrop filters, and an algorithmic Kinetic Word-Level Scroll Velocity Skew Hook (useScrollVelocitySkew). This custom physics engine computes real-time document scroll delta (Δy / Δt) and applies spring-damped geometric skew transformations to individual typographic elements on the fly.")
    add_bullet("• Multi-Tier Role-Based Access Control (RBAC) System:",
               "Built a secure, hierarchical authorization framework separating public users from authenticated personnel. Developed two distinct administrative suites: (1) AdminPage.jsx (160 KB) for day-to-day content updates, project curation, and contact log monitoring; and (2) SuperAdminPage.jsx (274 KB) for master controls including granular user access provisioning, real-time telemetry, audit logging, and one-click database snapshot exports (innoveity_cms_backup_database.json).")
    add_bullet("• Dual-Sync Database Orchestration Layer:",
               "Designed a unified data abstraction layer (CMSContext.jsx and mysqlClient.js) that automatically switches between Supabase Cloud PostgreSQL and MySQL 8. Implemented WebSocket subscriptions paired with a 3-second background polling heartbeat, ensuring instant global state synchronization whenever content is updated in the CMS.")
    add_bullet("• Reactive Automated SEO Suite:",
               "Engineered the proprietary SEOHead component using dynamic document head manipulation. The engine reactively updates meta titles, descriptions, canonical URLs, Twitter cards, and OpenGraph social metadata on every client-side route change, eliminating SEO drop-offs common in traditional Single Page Applications.")

    # ========================== 3. POTENTIAL USERS & STAKEHOLDERS ==========================
    add_h1("3. POTENTIAL USERS & STAKEHOLDERS")
    add_h2("Who are the target users and beneficiaries of Innoveity Tech Solution?")
    add_body_p(
        "The platform serves a multi-stakeholder ecosystem spanning enterprise leadership, institutional educators, students, and operational staff:"
    )

    add_num_bullet("1.", "Enterprise Clients & B2B Organizations (Corporate Buyers):",
                   "Mid-market enterprises, tech conglomerates, and funded startups seeking bespoke software engineering, enterprise resource planning (ERP) systems, cloud migration, AI agent integration, and mobile application development.")
    add_num_bullet("2.", "Engineering Colleges, Universities & Deans of Placement:",
                   "Higher education leaders across Tamil Nadu and India seeking structured corporate training partnerships, industry-aligned curriculum modernization, student hackathons, and verified campus recruitment drives.")
    add_num_bullet("3.", "Internal Marketing & Content Management Teams (Admins):",
                   "Non-technical internal staff who utilize the dedicated CMS to edit website copy, update statistics, showcase newly delivered client projects, publish press releases, and reorder executive profiles without editing raw code.")
    add_num_bullet("4.", "Corporate Recruiters & Talent Acquisition Partners:",
                   "Industry recruiters scouting pre-assessed, high-caliber engineering graduates trained through Innoveity’s intensive technical immersion programs.")
    add_num_bullet("5.", "Students, Engineering Aspirants & Working Professionals:",
                   "Class of graduating engineers and IT professionals seeking advanced training in Full-Stack Web Development, Cloud Computing, AI Engineering, and direct placement opportunities.")
    add_num_bullet("6.", "Executive Leadership & Platform Super Admins:",
                   "Innoveity’s executive board and system administrators who require real-time visibility into incoming project inquiries, database backups, audit records, and platform health telemetry.")

    # ========================== 4. COMPREHENSIVE SERVICE OFFERINGS ==========================
    add_h1("4. COMPREHENSIVE SERVICE OFFERINGS & DOMAIN PILLARS")
    add_h2("What core solutions are delivered through the Innoveity Tech Solution platform?")
    add_body_p("The platform orchestrates nine primary enterprise engineering pillars, each mapped to modular architectural components:")

    services_data = [
        ("01", "AI & Intelligent Solutions", "Custom Large Language Model (LLM) fine-tuning, voice AI interfaces, regional multilingual intelligence engines, sub-second natural language processing, and automated enterprise workflow AI agents."),
        ("02", "Enterprise Software Systems", "Mission-critical ERP, CRM, Management Information Systems (MIS), human resource platforms, automated compliance tracking, and role-based data architectures."),
        ("03", "Education Technology (EdTech)", "Proprietary Learning Management Systems (LMS), real-time student evaluation and automated assessment engines, digital placement portals, and comprehensive campus ERPs."),
        ("04", "Digital Governance & Public Sector Tech", "Citizen engagement applications, municipal service portals, real-time administrative intelligence dashboards, and tamper-resistant digital record management adhering to strict security protocols."),
        ("05", "Smart Mobility & Transit Solutions", "High-throughput QR-based and contactless ticketing gateways, automated commuter wayfinding applications, and real-time fleet telemetry."),
        ("06", "Logistics, Supply Chain & Field Tracking", "End-to-end parcel tracking pipelines, industrial barcode/QR scanning ecosystems, multi-depot inventory synchronizers, and field technician dispatch software."),
        ("07", "Web & Mobile Application Engineering", "Cross-platform mobile applications (React Native, iOS, Android), Progressive Web Apps (PWAs), high-concurrency microfrontends, and interactive corporate web platforms."),
        ("08", "IoT & Connected Smart Infrastructure", "Industrial IoT sensor integration, device gateway controllers, low-latency telemetry dashboards, and remote industrial equipment monitoring."),
        ("09", "Startup & Venture Innovation Lab", "Rapid Minimum Viable Product (MVP) prototyping, cloud-native foundational architecture design, and fractional CTO advisory services for early-stage and growth-stage technology ventures.")
    ]

    for num, name, desc in services_data:
        add_num_bullet(num + ".", name + ":", desc)

    # ========================== 5. STRUCTURED 9-STEP DELIVERY LIFECYCLE ==========================
    add_h1("5. STRUCTURED 9-STEP DELIVERY LIFECYCLE")
    add_h2("How does Innoveity Tech Solution execute client engineering programs?")
    add_body_p("Every engineering initiative follows a deterministic, 9-step delivery pipeline embedded within the platform's operational model:")

    process_steps = [
        ("01", "Ideation", "Deep requirements discovery, pain-point analysis, user persona mapping, and operational feasibility scoping."),
        ("02", "Strategy", "Formulation of technical roadmaps, project milestones, resource allocation, and risk mitigation matrices."),
        ("03", "UI/UX Design", "Wireframing, interactive Figma prototyping, visual design system creation, and ergonomic user experience validation."),
        ("04", "Architecture", "Selection of data models, database schemas, microservice topologies, and zero-trust security protocols."),
        ("05", "Development", "Agile sprint execution, modular clean-code development, version-controlled pull requests, and continuous peer review."),
        ("06", "AI Integration", "Embedding contextual intelligence, automated machine learning models, and predictive analytics into business workflows."),
        ("07", "Testing & QA", "Multi-device cross-browser validation, load testing, penetration security audits, and automated regression test suites."),
        ("08", "Deployment", "Containerized CI/CD release, edge caching configuration, serverless functions provisioning, and zero-downtime production rollouts."),
        ("09", "Support & Scaling", "Continuous SLA monitoring, performance tuning, database index optimization, and proactive feature enhancements.")
    ]

    lifecycle_table = doc.add_table(rows=len(process_steps) + 1, cols=3)
    lifecycle_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lifecycle_table.autofit = False
    set_table_borders(lifecycle_table, color="CBD5E1", sz="4", val="single")

    # Header Row
    hdr_cells = lifecycle_table.rows[0].cells
    hdr_titles = ["Phase", "Lifecycle Stage", "Operational Scope & Deliverables"]
    widths = [Inches(1.0), Inches(2.0), Inches(3.5)]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].width = widths[i]
        set_cell_shading(hdr_cells[i], "0F172A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = 'Segoe UI Semibold'
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (p_num, p_title, p_desc) in enumerate(process_steps):
        row = lifecycle_table.rows[idx + 1]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width, c1.width, c2.width = widths[0], widths[1], widths[2]
        
        bg_col = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        for c in (c0, c1, c2):
            set_cell_shading(c, bg_col)
            set_cell_margins(c, top=80, bottom=80, left=140, right=140)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(p_num)
        r0.font.name = 'Segoe UI Semibold'
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = BLUE_ACCENT

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(p_title)
        r1.font.name = 'Segoe UI Semibold'
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.color.rgb = NAVY

        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(p_desc)
        r2.font.name = 'Calibri'
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = BODY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ========================== 6. UNIQUE FEATURES & INNOVATIONS ==========================
    add_h1("6. UNIQUE FEATURES & TECHNICAL INNOVATIONS")
    add_h2("What sets Innoveity Tech Solution apart from legacy agency portals and conventional CMS software?")

    add_bullet("• Zero-Rebuild Instant CMS Propagation:",
               "Traditional headless platforms (e.g., Strapi or WordPress connected to Next.js/Gatsby) require full static site rebuilds or CDN cache revalidation when copy changes. Innoveity Tech utilizes real-time client state hydration powered by WebSockets, allowing updates made in the admin panel to reflect across all active browser windows worldwide in under 500 milliseconds.")
    add_bullet("• Algorithmic Kinetic Typography & Scroll Physics:",
               "Replaces static CSS animations with an active mathematical physics loop (useScrollVelocitySkew). As users scroll through the website, words calculate their individual velocity vector and tilt smoothly along their horizontal axis with spring-damped elasticity, creating a fluid, living digital experience.")
    add_bullet("• Active-Passive Dual-Database Resilience:",
               "The platform eliminates vendor lock-in and single-point-of-failure risks. It runs on a dual-database backbone: Supabase Cloud PostgreSQL handles real-time cloud transactions, while an on-premise/cloud MySQL 8 server with an automated seeding pipeline (server/seed.js) provides reliable persistence and offline recovery.")
    add_bullet("• Two-Tier Administrative Separation of Concerns:",
               "Instead of exposing high-risk configuration to regular staff, the system enforces a strict dual-dashboard model: The Admin Panel focuses exclusively on operational tasks (project listings, media cards, team roster, contact logs), while the Super Admin Panel encapsulates master controls (database export/import, role permissions, system diagnostics, and administrative credential management).")
    add_bullet("• Automated SPA-Aware Dynamic SEO Suite:",
               "Single-page React apps historically struggle with dynamic metadata indexing. Innoveity Tech integrates the SEOHead micro-engine, which intercepts React Router v7 navigation events and dynamically regenerates document title tags, schema JSON-LD, OpenGraph images, and canonical URLs before web crawlers process the page.")
    add_bullet("• Direct Dual-Channel Communication Engine:",
               "Inquiry forms run on a dual-routing pipeline: primary delivery via browser-level EmailJS dispatch, followed by automated fallback to an Express.js SMTP relay (via Gmail/cPanel SMTP), ensuring critical enterprise inquiries never bounce.")

    # ========================== 7. UNIQUE SELLING PROPOSITION (USP) ==========================
    add_h1("7. UNIQUE SELLING PROPOSITION (USP)")
    add_h2("Why should enterprise clients and institutions partner with Innoveity Tech Solution?")

    add_num_bullet("1.", "Dual-Impact Value Engine (Enterprise Engineering + Academic Talent Pipeline):",
                   "Most technology agencies solely write code; educational consultancies solely provide training. Innoveity Tech is uniquely positioned at the intersection of both. We build mission-critical enterprise platforms while simultaneously training top engineering cohorts across 50+ universities. Clients benefit from state-of-the-art software delivery and direct recruitment access to pre-trained, high-caliber software developers.")
    add_num_bullet("2.", "Proprietary, High-Performance Codebase (Zero WordPress Bloat):",
                   "Innoveity Tech completely eschews third-party PHP plugins, bloated WordPress themes, and high-vulnerability extensions. The entire platform is built with hand-crafted React 19, Vite, and modular CSS tokens, yielding ultra-high Lighthouse scores, instantaneous interactions, and bulletproof security.")
    add_num_bullet("3.", "Complete Autonomous Content Control Without Retainer Costs:",
                   "Clients are freed from recurring developer fees for basic updates. Through the bespoke Admin CMS, marketing and HR teams can update testimonials, project case studies, leadership profiles, and corporate contact details in seconds with zero coding knowledge.")
    add_num_bullet("4.", "Proven Institutional Track Record & Metric-Driven Outcomes:",
                   "Innoveity’s credibility is backed by verified academic results: 5+ Years of Industry Leadership (Founded in 2019), 500+ Projects Delivered across Web, Mobile, Cloud, and AI, 50+ Tech Experts in active development and research, 99% Client Retention and Project Success Rate, 85% Average Placement Improvement for partner engineering colleges, and 2,000+ Students Trained Annually in core and modern technical stacks.")
    add_num_bullet("5.", "Strategic Innovation Hub:",
                   "Headquartered inside the prestigious MCC MRF Innovation Park in East Tambaram, Chennai, the organization operates in direct proximity to Tamil Nadu’s premier academic corridors and technology talent pools, facilitating rapid deployment and client collaboration.")

    # ========================== 8. SYSTEM ARCHITECTURE SPECIFICATION ==========================
    add_h1("8. SYSTEM ARCHITECTURE SPECIFICATION")
    add_h2("Technical Topology, Data Flow & Component Taxonomy")
    add_body_p("The Innoveity Tech Solution platform is structured into four decoupled, fault-tolerant tiers:")

    arch_points = [
        ("Tier 1: Presentation & Motion Engine (Client Layer)", "Built on React 19, Vite 8, and React Router v7. Features Framer Motion for entrance choreography and the proprietary useScrollVelocitySkew hook for kinetic word tilting. UI styling relies on vanilla CSS custom properties and HSL tokens without framework overhead."),
        ("Tier 2: Dual Application Portals (RBAC Layer)", "Separates public browsing routes from administrative endpoints via the AdminRouteDispatcher. Contains the 160 KB AdminPage.jsx for standard CMS curation and the 274 KB SuperAdminPage.jsx for user authentication, role management, and database snapshot exports."),
        ("Tier 3: State & Dual-Sync Engine (Abstraction Layer)", "Orchestrated by CMSContext.jsx. Maintains an in-memory cache synchronized in real time via Supabase WebSocket listeners, backed by a 3-second background polling cycle against the Express.js / MySQL API as a persistent failover."),
        ("Tier 4: Enterprise Persistence & Communication Tier", "Primary real-time data hosted on Supabase Cloud PostgreSQL. Relational persistence provided by MySQL 8 (via mysql2 connection pooling). Outbound inquiries routed via browser EmailJS with automated fallback to an Express.js SMTP relay.")
    ]

    for t_name, t_desc in arch_points:
        add_bullet("• " + t_name + ":", t_desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Document Signoff / Certification Box
    signoff_table = doc.add_table(rows=1, cols=1)
    signoff_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = signoff_table.rows[0].cells[0]
    c.width = Inches(6.5)
    set_cell_shading(c, "F1F5F9")
    set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    set_table_borders(signoff_table, color="94A3B8", sz="6", val="single")

    sp = c.paragraphs[0]
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(2)
    srun1 = sp.add_run("DOCUMENT CERTIFICATION & COPYRIGHT NOTICE\n")
    srun1.font.name = 'Segoe UI Semibold'
    srun1.font.size = Pt(10.5)
    srun1.font.bold = True
    srun1.font.color.rgb = NAVY

    srun2 = sp.add_run(
        "© 2026 Innoveity Tech Solution Private Limited. All Rights Reserved.\n"
        "Conceptualized, Architected, and Engineered by Lead Architect: Nancy Narmadha T.\n"
        "Headquarters: MCC MRF Innovation Park, East Tambaram, Chennai - 600059, Tamil Nadu, India.\n"
        "Official Website: https://innoveitytech.com  |  Direct Contact: innoveitytech@gmail.com | +91 7904327211"
    )
    srun2.font.name = 'Calibri'
    srun2.font.size = Pt(10)
    srun2.font.color.rgb = MUTED_COLOR

    output_path = r"n:\Innoveity-Tech--main\INNOVEITY_TECH_PRODUCT_OVERVIEW.docx"
    doc.save(output_path)
    print(f"Successfully generated executive Word document at: {output_path}")

if __name__ == "__main__":
    build_document()
